"""Make anonymized copies of the weekly Excel reports.

The simulation reports name the team's cafe and the other teams' cafes. Before the data
could go on GitHub those names had to go, without touching any of the numbers.

    python scripts/anonymize_workbooks.py SOURCE_DIR DEST_DIR --mapping names.json

`names.json` maps each real name to its replacement, for example
{"Real Cafe Name": "Cafe A", "Other Team": "Cafe B"}. The mapping file itself is never
committed. The originals are only read; sanitized copies are written to DEST_DIR with the
same "Week N/<report>.xlsx" layout.

What gets cleaned:
  - every text cell (exact names and names inside longer text such as "Real Cafe - Week 3")
  - sheet names
  - cell comments
  - workbook properties (author, last modified by, title, ...)
Numeric cells and their number formats are left exactly as they are.
"""
from __future__ import annotations

import argparse
import json
import re
import shutil
import sys
from pathlib import Path

import openpyxl

WEEK_FOLDER = re.compile(r"^week\s*\d+$", re.IGNORECASE)


def build_pattern(mapping: dict[str, str]) -> re.Pattern[str]:
    # Longest names first so "Blue Cafe Two" is matched before "Blue Cafe".
    names = sorted(mapping, key=len, reverse=True)
    return re.compile("|".join(re.escape(n) for n in names))


def replace_names(text: str, mapping: dict[str, str], pattern: re.Pattern[str] | None = None) -> str:
    """Replace every mapped name inside `text`. Non-matching text is returned unchanged."""
    if not mapping:
        return text
    pattern = pattern or build_pattern(mapping)
    return pattern.sub(lambda m: mapping[m.group(0)], text)


def anonymize_workbook(src: Path, dest: Path, mapping: dict[str, str]) -> dict[str, int]:
    """Copy one workbook, replacing names everywhere text can live. Returns change counts."""
    pattern = build_pattern(mapping) if mapping else None
    counts = {"cells": 0, "sheets": 0, "comments": 0, "properties": 0}
    wb = openpyxl.load_workbook(src)
    for ws in wb.worksheets:
        new_title = replace_names(ws.title, mapping, pattern)
        if new_title != ws.title:
            ws.title = new_title
            counts["sheets"] += 1
        for row in ws.iter_rows():
            for cell in row:
                if isinstance(cell.value, str):
                    new = replace_names(cell.value, mapping, pattern)
                    if new != cell.value:
                        cell.value = new
                        counts["cells"] += 1
                if cell.comment is not None:
                    # Comments carry an author field as well as text; drop them entirely.
                    cell.comment = None
                    counts["comments"] += 1
    props = wb.properties
    for attr in ("creator", "lastModifiedBy", "title", "subject", "description",
                 "keywords", "category", "company", "manager"):
        if getattr(props, attr, None):
            setattr(props, attr, None)
            counts["properties"] += 1
    dest.parent.mkdir(parents=True, exist_ok=True)
    wb.save(dest)
    wb.close()
    return counts


def anonymize_folder(source: Path, dest: Path, mapping: dict[str, str]) -> list[tuple[str, dict[str, int]]]:
    """Process every workbook inside the "Week N" folders of `source`."""
    results = []
    for week_dir in sorted(p for p in source.iterdir() if p.is_dir() and WEEK_FOLDER.match(p.name)):
        for wb_path in sorted(week_dir.glob("*.xlsx")):
            if wb_path.name.startswith("~$"):
                continue  # Excel lock files
            target = dest / week_dir.name / wb_path.name
            results.append((f"{week_dir.name}/{wb_path.name}", anonymize_workbook(wb_path, target, mapping)))
    return results


def find_remaining(dest: Path, names: list[str]) -> list[tuple[str, str, str]]:
    """Scan the sanitized copies for any name that slipped through."""
    hits = []
    for wb_path in sorted(dest.rglob("*.xlsx")):
        wb = openpyxl.load_workbook(wb_path, read_only=True)
        for ws in wb.worksheets:
            for row in ws.iter_rows():
                for cell in row:
                    if isinstance(cell.value, str) and any(n in cell.value for n in names):
                        hits.append((str(wb_path.relative_to(dest)), ws.title, cell.coordinate))
        wb.close()
    return hits


def main(argv: list[str] | None = None) -> int:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("source", type=Path, help="folder containing the original Week N folders")
    ap.add_argument("dest", type=Path, help="where to write the sanitized copies")
    ap.add_argument("--mapping", type=Path, required=True, help="JSON file: {real name: replacement}")
    ap.add_argument("--clean", action="store_true", help="delete DEST first")
    args = ap.parse_args(argv)

    mapping = json.loads(args.mapping.read_text(encoding="utf-8"))
    if args.clean and args.dest.exists():
        shutil.rmtree(args.dest)
    results = anonymize_folder(args.source, args.dest, mapping)
    for name, counts in results:
        changed = {k: v for k, v in counts.items() if v}
        print(f"{name}: {changed or 'nothing to change'}")
    leftovers = find_remaining(args.dest, list(mapping))
    if leftovers:
        print(f"\nWARNING: {len(leftovers)} cell(s) still contain an original name:", file=sys.stderr)
        for hit in leftovers:
            print("  ", hit, file=sys.stderr)
        return 1
    print(f"\n{len(results)} workbooks written to {args.dest}; no original names remain.")
    return 0


if __name__ == "__main__":
    sys.exit(main())
