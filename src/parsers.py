"""Parsers for the weekly simulation reports.

Each report is an Excel workbook laid out for people to read, not for software: a title
block, then one or more small tables separated by blank rows, with labels in the first
column. Different reports have different layouts, and the same report can shift by a row
or two from week to week, so the parsers never rely on fixed cell addresses. They look
for the labels ("Measure", "Company News", "Comments/Suggestions", ...) and read the rows
relative to them.

Every parsed fact keeps the file, sheet and cell it came from, so any number Lumo quotes
can be traced back to the spreadsheet.
"""
from __future__ import annotations

import re
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any, Iterator

import openpyxl
from openpyxl.utils import get_column_letter

from .values import ParsedValue, parse_value

TITLE_RE = re.compile(r"^(?P<company>.+?)\s*-\s*Week\s+(?P<week>\d+)\s*$", re.IGNORECASE)
PERIOD_RE = re.compile(r"(January|February|March|April|May|June|July|August|September|October|November|December)"
                       r"\s*,\s*(Month to Date|End of Month|Year to Date)", re.IGNORECASE)


def clean(text: Any) -> str:
    """Collapse whitespace in a label; keep the original case."""
    return re.sub(r"\s+", " ", str(text)).strip() if text is not None else ""


def norm(text: Any) -> str:
    """Lower-case, whitespace-collapsed label for comparisons."""
    return clean(text).lower()


# --------------------------------------------------------------------------- worksheet grid
@dataclass(frozen=True)
class Cell:
    row: int
    col: int
    value: Any
    number_format: str | None

    @property
    def coordinate(self) -> str:
        return f"{get_column_letter(self.col)}{self.row}"

    @property
    def is_blank(self) -> bool:
        return self.value is None or (isinstance(self.value, str) and not self.value.strip())

    def parsed(self) -> ParsedValue:
        return parse_value(self.value, self.number_format)


class Grid:
    """A worksheet as a sparse dictionary of non-empty cells, with label-search helpers."""

    def __init__(self, worksheet) -> None:
        self.title: str = worksheet.title
        self.cells: dict[tuple[int, int], Cell] = {}
        self.max_row = 0
        for row in worksheet.iter_rows():
            for c in row:
                if c.value is not None:
                    self.cells[(c.row, c.column)] = Cell(c.row, c.column, c.value, c.number_format)
                    self.max_row = max(self.max_row, c.row)

    def cell(self, row: int, col: int) -> Cell:
        return self.cells.get((row, col), Cell(row, col, None, None))

    def text(self, row: int, col: int) -> str:
        return clean(self.cell(row, col).value)

    def row_cells(self, row: int) -> list[Cell]:
        return [c for (r, _), c in sorted(self.cells.items()) if r == row]

    def is_blank_row(self, row: int) -> bool:
        return all(c.is_blank for c in self.row_cells(row))

    def find(self, label: str) -> Cell | None:
        """First cell whose text equals `label` (case-insensitive)."""
        target = norm(label)
        for (_, _), c in sorted(self.cells.items()):
            if isinstance(c.value, str) and norm(c.value) == target:
                return c
        return None

    def find_header_row(self, required: list[str]) -> tuple[int, dict[str, int]] | None:
        """First row that contains all `required` labels. Returns (row, {label: column})."""
        wanted = {norm(x) for x in required}
        for r in range(1, self.max_row + 1):
            headers = {norm(c.value): c.col for c in self.row_cells(r) if isinstance(c.value, str)}
            if wanted <= headers.keys():
                return r, headers
        return None

    def table_rows(self, start_row: int) -> Iterator[int]:
        """Row numbers from `start_row` until the first blank row."""
        r = start_row
        while r <= self.max_row and not self.is_blank_row(r):
            yield r
            r += 1

    def title_info(self) -> dict[str, Any]:
        """Company and week from the title block ("Cafe A - Week 7"), if present."""
        for r in range(1, min(self.max_row, 4) + 1):
            for c in self.row_cells(r):
                if isinstance(c.value, str):
                    m = TITLE_RE.match(clean(c.value))
                    if m:
                        return {"company": clean(m.group("company")), "week": int(m.group("week"))}
        return {"company": None, "week": None}


def load_grids(path: str | Path) -> list[Grid]:
    workbook = openpyxl.load_workbook(path, data_only=True)
    try:
        return [Grid(ws) for ws in workbook.worksheets]
    finally:
        workbook.close()


# --------------------------------------------------------------------------- parse result
@dataclass
class ParsedReport:
    week: int
    source_file: str
    report_type: str
    company: str | None = None
    metrics: list[dict[str, Any]] = field(default_factory=list)
    comments: list[dict[str, Any]] = field(default_factory=list)
    text_records: list[dict[str, Any]] = field(default_factory=list)
    transactions: list[dict[str, Any]] = field(default_factory=list)
    daily_receipts: list[dict[str, Any]] = field(default_factory=list)
    decisions: list[dict[str, Any]] = field(default_factory=list)
    warnings: list[str] = field(default_factory=list)

    def where(self, sheet: str, cell: str | None) -> dict[str, Any]:
        return {"week": self.week, "source_file": self.source_file, "sheet_name": sheet, "cell": cell}

    def add_metric(self, section: str, metric: str, parsed: ParsedValue, sheet: str, cell: str, *,
                   company: str | None = None, rank: int | None = None, change: float | None = None,
                   period: str | None = None, period_label: str | None = None, line_order: int | None = None,
                   unit: str | None = None) -> None:
        self.metrics.append({
            **self.where(sheet, cell), "report_type": self.report_type, "section": section,
            "company": company if company is not None else self.company, "metric": clean(metric),
            "value_numeric": parsed.numeric, "value_text": parsed.text, "unit": unit or parsed.unit,
            "scale": parsed.scale, "qualifier": parsed.qualifier, "rank": rank, "change_reported": change,
            "period": period or parsed.period, "period_label": period_label, "line_order": line_order,
        })

    def add_text(self, section: str, text: str, sheet: str, cell: str, order: int) -> None:
        self.text_records.append({**self.where(sheet, cell), "report_type": self.report_type,
                                  "section": section, "text": clean(text), "text_order": order})


# --------------------------------------------------------------------------- dashboard
def parse_dashboard(grids: list[Grid], report: ParsedReport) -> None:
    """Measure | Value | Change | Rank table, then Industry News and Company News lists."""
    grid = grids[0]
    header = grid.find_header_row(["Measure", "Value"])
    if header is None:
        report.warnings.append("dashboard: 'Measure' header row not found")
        return
    hdr_row, cols = header
    for order, r in enumerate(grid.table_rows(hdr_row + 1)):
        label = grid.cell(r, cols["measure"])
        if not isinstance(label.value, str):
            continue
        value = grid.cell(r, cols["value"])
        change = grid.cell(r, cols["change"]).parsed().numeric if "change" in cols else None
        rank_value = grid.cell(r, cols["rank"]).parsed().numeric if "rank" in cols else None
        report.add_metric("Measures", label.value, value.parsed(), grid.title, value.coordinate,
                          rank=int(rank_value) if rank_value is not None else None, change=change, line_order=order)

    for section in ("Industry News", "Company News"):
        start = grid.find(section)
        if start is None:
            continue
        for order, r in enumerate(grid.table_rows(start.row + 1)):
            c = grid.cell(r, start.col)
            if norm(c.value) in ("industry news", "company news"):
                break  # the next section can follow without a blank row
            if not c.is_blank:
                report.add_text(section, str(c.value), grid.title, c.coordinate, order)


# --------------------------------------------------------------------------- survey
def parse_survey(grids: list[Grid], report: ParsedReport) -> None:
    """Company satisfaction (us vs industry), category ratings, and weighted comments."""
    grid = grids[0]
    header = grid.find_header_row(["Company", "Satisfaction"])
    if header:
        hdr_row, cols = header
        for order, r in enumerate(grid.table_rows(hdr_row + 1)):
            name = grid.cell(r, cols["company"])
            if name.is_blank:
                continue
            # "Industry Average (11 Companies)" -> "Industry Average"
            company = re.sub(r"\s*\(.*?\)\s*$", "", clean(name.value))
            if order == 0 and not company.lower().startswith("industry"):
                report.company = company  # the first row is the reporting cafe
            value = grid.cell(r, cols["satisfaction"])
            change = grid.cell(r, cols["change"]).parsed().numeric if "change" in cols else None
            report.add_metric("Company Satisfaction", "Satisfaction", value.parsed(), grid.title, value.coordinate,
                              company=company, change=change, line_order=order, unit="percent")

    header = grid.find_header_row(["Survey Category", "Rating"])
    if header:
        hdr_row, cols = header
        for order, r in enumerate(grid.table_rows(hdr_row + 1)):
            category = grid.cell(r, cols["survey category"])
            value = grid.cell(r, cols["rating"])
            if category.is_blank:
                continue
            report.add_metric("Survey Category Ratings", f"{clean(category.value)} Rating", value.parsed(),
                              grid.title, value.coordinate, line_order=order, unit="rating")

    header = grid.find_header_row(["Comments/Suggestions", "Weight"])
    if header:
        hdr_row, cols = header
        for order, r in enumerate(grid.table_rows(hdr_row + 1)):
            text = grid.cell(r, cols["comments/suggestions"])
            if text.is_blank:
                continue
            weight = grid.cell(r, cols["weight"]).parsed().numeric
            report.comments.append({**report.where(grid.title, text.coordinate), "company": report.company,
                                    "comment": clean(text.value), "weight": weight, "comment_order": order})


# --------------------------------------------------------------------------- labor
def parse_labor(grids: list[Grid], report: ParsedReport) -> None:
    """Local labor market averages ("$931.25 / week") and employees per company."""
    grid = grids[0]
    header = grid.find_header_row(["Measure", "Local Average"])
    if header:
        hdr_row, cols = header
        for order, r in enumerate(grid.table_rows(hdr_row + 1)):
            label, value = grid.cell(r, cols["measure"]), grid.cell(r, cols["local average"])
            if not label.is_blank:
                report.add_metric("Local Labor Market", label.value, value.parsed(), grid.title, value.coordinate,
                                  company=None, line_order=order)

    header = grid.find_header_row(["Company", "Employees"])
    if header:
        hdr_row, cols = header
        for order, r in enumerate(grid.table_rows(hdr_row + 1)):
            name, value = grid.cell(r, cols["company"]), grid.cell(r, cols["employees"])
            if name.is_blank:
                continue
            company = clean(name.value)
            if isinstance(name.value, float) and name.value.is_integer():
                company = str(int(name.value))  # unused team slots are numbered, keep them as text
            report.add_metric("Employees by Company", "Employees", value.parsed(), grid.title, value.coordinate,
                              company=company, line_order=order, unit="count")


# --------------------------------------------------------------------------- receipts
DAILY_COLUMNS = {
    "daily capacity": "daily_capacity", "cups served": "cups_served", "avg. price": "avg_price",
    "receipts": "receipts", "satisfied": "satisfied", "long wait": "long_wait",
    "served after hours": "served_after_hours", "left or outside hours": "left_or_outside",
}


def parse_receipts(grids: list[Grid], report: ParsedReport) -> None:
    """Daily traffic rows (+ a Total row) and the weekly product sales table."""
    grid = grids[0]
    header = grid.find_header_row(["Date", "Day", "Cups Served", "Receipts"])
    if header is None:
        report.warnings.append("receipts: daily header row not found")
        return
    hdr_row, cols = header
    header_labels = {c.col: clean(c.value) for c in grid.row_cells(hdr_row)}
    for order, r in enumerate(grid.table_rows(hdr_row + 1)):
        date = grid.cell(r, cols["date"])
        is_total = norm(date.value) == "total"
        row: dict[str, Any] = {**report.where(grid.title, f"A{r}"), "row_order": order,
                               "date_text": clean(date.value) or None, "day": grid.text(r, cols["day"]) or None,
                               "is_total": is_total, "left_or_outside_text": None}
        for label, field_name in DAILY_COLUMNS.items():
            value = grid.cell(r, cols[label]) if label in cols else None
            parsed = value.parsed() if value else ParsedValue(None, "")
            row[field_name] = parsed.numeric
            if field_name == "left_or_outside":
                row["left_or_outside_text"] = parsed.text or None
            if is_total and value is not None and not value.is_blank:
                # The Total row doubles as the week's traffic metrics ("Cups Served", "Long Wait", ...)
                report.add_metric("Weekly Traffic Totals", header_labels[value.col], parsed, grid.title, value.coordinate)
        report.daily_receipts.append(row)

    totals = grid.find("Weekly Totals")
    if totals is None:
        report.warnings.append("receipts: no 'Weekly Totals' product table in this week")
        return
    cols = {norm(c.value): c.col for c in grid.row_cells(totals.row) if isinstance(c.value, str)}
    for order, r in enumerate(grid.table_rows(totals.row + 1)):
        name = grid.cell(r, totals.col)
        if name.is_blank:
            continue
        product = clean(name.value)
        for label, suffix in (("units", "Units"), ("unit price", "Unit Price"), ("receipts", "Receipts")):
            if label not in cols:
                continue
            value = grid.cell(r, cols[label])
            if value.is_blank:
                continue
            metric = "Total Receipts" if norm(product) == "total" else f"{product} {suffix}"
            report.add_metric("Product Sales", metric, value.parsed(), grid.title, value.coordinate, line_order=order)


# --------------------------------------------------------------------------- checkbook
def parse_checkbook(grids: list[Grid], report: ParsedReport) -> None:
    """Every transaction, plus derived beginning/ending balance and totals per account."""
    grid = grids[0]
    header = grid.find_header_row(["Date", "Description", "Account", "Balance"])
    if header is None:
        report.warnings.append("checkbook: header row not found")
        return
    hdr_row, cols = header
    totals: dict[str, dict[str, float]] = {}
    last_balance: Cell | None = None
    order = 0
    for r in grid.table_rows(hdr_row + 1):
        description = grid.text(r, cols["description"])
        balance = grid.cell(r, cols["balance"])
        if norm(description) == "beginning balance":
            report.add_metric("Checkbook Balance", "Beginning Balance", balance.parsed(), grid.title,
                              balance.coordinate, unit="USD")
            continue
        payment = grid.cell(r, cols["payment"]).parsed().numeric if "payment" in cols else None
        deposit = grid.cell(r, cols["deposit"]).parsed().numeric if "deposit" in cols else None
        account = grid.text(r, cols["account"]) or None
        report.transactions.append({**report.where(grid.title, f"A{r}"), "txn_order": order,
                                    "date_text": grid.text(r, cols["date"]) or None, "description": description or None,
                                    "account": account, "payment": payment, "deposit": deposit,
                                    "balance": balance.parsed().numeric})
        if account:
            acc = totals.setdefault(account, {"Payments": 0.0, "Deposits": 0.0})
            acc["Payments"] += payment or 0.0
            acc["Deposits"] += deposit or 0.0
        last_balance = balance
        order += 1
    if last_balance is not None:
        report.add_metric("Checkbook Balance", "Ending Balance", last_balance.parsed(), grid.title,
                          last_balance.coordinate, unit="USD")
    for i, (account, sums) in enumerate(totals.items()):
        for kind, amount in sums.items():
            if amount:
                parsed = ParsedValue(round(amount, 2), f"sum of {kind.lower()} for {account}", unit="USD")
                report.add_metric("Checkbook Account Totals", f"{account} {kind}", parsed, grid.title,
                                  f"derived:{account}", line_order=i)


# --------------------------------------------------------------------------- financial statements
def parse_financial_statement(grids: list[Grid], report: ParsedReport) -> None:
    """Income statement, balance sheet or cash flow statement.

    Only the first sheet (the statement itself) is parsed; the workbook's other sheets are
    per-account transaction listings that the checkbook already covers. Statement values
    are cumulative for the month (or year) to date, not amounts for the week, so each row
    records its period.
    """
    grid = grids[0]
    label_cell = next((c for r in range(1, 4) for c in grid.row_cells(r)
                       if isinstance(c.value, str) and PERIOD_RE.search(c.value)), None)
    period_label = None
    if label_cell:
        m = PERIOD_RE.search(label_cell.value)
        period_label = f"{m.group(1).title()}, {m.group(2)}"

    # Income statements have two value columns (month, year to date) separated by "% of Rev."
    header = grid.find_header_row(["% of Rev."])
    value_columns: list[tuple[int, str]] = []
    first_row = 3
    if header:
        hdr_row, _ = header
        first_row = hdr_row + 1
        for c in grid.row_cells(hdr_row):
            if isinstance(c.value, str) and "%" not in c.value:
                value_columns.append((c.col, "year_to_date" if "year" in norm(c.value) else "month_to_date"))
    else:
        value_columns = [(2, "month_to_date")]

    section = grid.title
    order = 0
    for r in range(first_row, grid.max_row + 1):
        label = grid.cell(r, 1)
        if not isinstance(label.value, str):
            continue
        first_value = grid.cell(r, value_columns[0][0])
        if first_value.is_blank or isinstance(first_value.value, str):
            section = clean(label.value)  # a label with no amount is a section header ("Assets")
            continue
        for col, period in value_columns:
            value = grid.cell(r, col)
            if value.is_blank:
                continue
            report.add_metric(section, label.value, value.parsed(), grid.title, value.coordinate,
                              period=period, period_label=period_label, line_order=order)
        order += 1


# --------------------------------------------------------------------------- decisions review
def parse_decisions(grids: list[Grid], report: ParsedReport) -> None:
    """Decision summary: one column per week, label-only rows act as section headers."""
    grid = grids[0]
    week_cols: dict[int, int] = {}
    hdr_row = None
    for r in range(1, min(grid.max_row, 8) + 1):
        found = {c.col: int(m.group(1)) for c in grid.row_cells(r)
                 if isinstance(c.value, str) and (m := re.match(r"^Week\s+(\d+)$", clean(c.value), re.IGNORECASE))}
        if found:
            week_cols, hdr_row = found, r
            break
    if hdr_row is None:
        report.warnings.append("decisions: no 'Week N' column headers found")
        return
    section = None
    for r in range(hdr_row + 1, grid.max_row + 1):
        label = grid.cell(r, 1)
        if not isinstance(label.value, str):
            continue
        values = [(col, grid.cell(r, col)) for col in week_cols if not grid.cell(r, col).is_blank]
        if not values:
            section = clean(label.value)
            continue
        for col, cell in values:
            parsed = cell.parsed()
            report.decisions.append({**report.where(grid.title, cell.coordinate), "decision_week": week_cols[col],
                                     "section": section, "item": clean(label.value).rstrip(":"),
                                     "value_text": parsed.text, "value_numeric": parsed.numeric})


# --------------------------------------------------------------------------- registry
REPORT_TYPES: dict[str, tuple[str, Any]] = {
    "results-dashboard": ("dashboard", parse_dashboard),
    "market-survey": ("survey", parse_survey),
    "market-labor": ("labor", parse_labor),
    "results-receipts": ("receipts", parse_receipts),
    "results-checkbook": ("checkbook", parse_checkbook),
    "results-income_statement": ("income_statement", parse_financial_statement),
    "results-balance_sheet": ("balance_sheet", parse_financial_statement),
    "results-cash_flow": ("cash_flow", parse_financial_statement),
    "decisions-review": ("decisions", parse_decisions),
}


def report_type_for(filename: str) -> str | None:
    stem = Path(filename).stem.lower()
    for pattern, (report_type, _) in REPORT_TYPES.items():
        if pattern in stem:
            return report_type
    return None


def parse_report(path: str | Path, week: int, source_file: str) -> ParsedReport:
    """Parse one workbook into a ParsedReport. Raises ValueError for unknown report types."""
    path = Path(path)
    report_type = report_type_for(path.name)
    if report_type is None:
        raise ValueError(f"no parser for {path.name}")
    parser = next(fn for rt, fn in REPORT_TYPES.values() if rt == report_type)
    grids = load_grids(path)
    report = ParsedReport(week=week, source_file=source_file, report_type=report_type)
    info = grids[0].title_info()
    report.company = info["company"]
    if info["week"] is not None and info["week"] != week:
        report.warnings.append(f"folder says week {week} but the report title says week {info['week']}")
    parser(grids, report)
    return report
